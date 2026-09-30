#!/usr/bin/env python3
"""Reproduce Figure 6 endpoint/missingness audit from the pinned public workbook.

Dependency: openpyxl. No cross-sheet subject linkage, fitted treatment model,
imputation, or raw-flow gating is performed. Run from any working directory.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import statistics
import urllib.request
from collections import Counter
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SOURCE_NAME = "41467_2026_72746_MOESM4_ESM.xlsx"
SOURCE_URL = "https://media.springernature.com/original/springer-static/esm/art%3A10.1038%2Fs41467-026-72746-5/MediaObjects/" + SOURCE_NAME
SOURCE_SHA256 = "3e302d90157981d46c04b72bc94633ad2fdf55381f016a33455d3986f35589e6"
ARTICLE_URL = "https://www.nature.com/articles/s41467-026-72746-5"
HORIZONS = [20, 40, 60, 120, 399]
GROUPS = [
    {"id": "pbs_vehicle", "label": "PBS + Vehicle", "survival_column": "C", "bli_start": 2, "bli_end": 11, "expected_n": 10},
    {"id": "pbs_tmz", "label": "PBS + TMZ", "survival_column": "D", "bli_start": 12, "bli_end": 20, "expected_n": 9},
    {"id": "rli_vehicle", "label": "RRV-RLI + Vehicle", "survival_column": "E", "bli_start": 21, "bli_end": 29, "expected_n": 9},
    {"id": "rli_tmz", "label": "RRV-RLI + TMZ", "survival_column": "F", "bli_start": 30, "bli_end": 38, "expected_n": 9},
]


def clean(value):
    return " ".join(str(value).replace("\xa0", " ").split())


def is_number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def km(records, mode):
    """Right-continuous Kaplan-Meier; events precede censors at tied times."""
    event_key = "paper_event" if mode == "paper_compatible" else "composite_event"
    n = len(records)
    survival = 1.0
    points = [{"time": 0, "survival": 1.0, "at_risk": n, "events": 0, "censored": 0, "remaining_after": n}]
    for t in sorted({r["time"] for r in records}):
        at_risk = sum(r["time"] >= t for r in records)
        events = sum(r["time"] == t and r[event_key] for r in records)
        censored = sum(r["time"] == t and not r[event_key] for r in records)
        survival *= 1 - events / at_risk
        points.append({"time": t, "survival": survival, "at_risk": at_risk, "events": events, "censored": censored, "remaining_after": at_risk - events - censored})
    return points


def rmst(records, mode, tau):
    points = km(records, mode)
    # Beyond the last observation a positive survival tail is not identifiable.
    # A tail already reaching zero contributes exactly zero thereafter.
    if tau > max(r["time"] for r in records) and points[-1]["survival"] > 1e-12:
        return None
    area, previous_t, previous_s = 0.0, 0.0, 1.0
    for point in points[1:]:
        t = point["time"]
        if t >= tau:
            return area + (tau - previous_t) * previous_s
        area += (t - previous_t) * previous_s
        previous_t, previous_s = t, point["survival"]
    return area + (tau - previous_t) * previous_s


def csv_write(path, rows, fields):
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def json_write(path, obj):
    path.write_text(json.dumps(obj, indent=2, ensure_ascii=False, allow_nan=False) + "\n", encoding="utf-8")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=ROOT / "research-inputs" / SOURCE_NAME)
    args = parser.parse_args()
    if not args.source.exists():
        args.source.parent.mkdir(parents=True, exist_ok=True)
        with urllib.request.urlopen(SOURCE_URL, timeout=60) as response:
            payload = response.read()
        if hashlib.sha256(payload).hexdigest() != SOURCE_SHA256:
            raise ValueError("Downloaded source checksum differs; do not silently update provenance")
        args.source.write_bytes(payload)
    actual_hash = hashlib.sha256(args.source.read_bytes()).hexdigest()
    if actual_hash != SOURCE_SHA256:
        raise ValueError(f"Source checksum mismatch: {actual_hash}")
    workbook = openpyxl.load_workbook(args.source, data_only=True)
    source_survival, source_bli = workbook["Figure 6B"], workbook["Figure 6C-G"]
    records = []
    for group in GROUPS:
        assert clean(source_survival[f'{group["survival_column"]}1'].value) == group["label"]
        for row in range(2, source_survival.max_row + 1):
            cell = source_survival[f'{group["survival_column"]}{row}']
            if cell.value is None:
                continue
            code = clean(cell.value)
            assert code in {"0", "1", "0*"}, (cell.coordinate, code)
            t = source_survival[f"A{row}"].value
            assert is_number(t) and t > 0
            records.append({
                "id": f"fig6B-r{row}", "group": group["id"], "time": t,
                "source_code": code,
                "event_type": {"1": "tumor_endpoint", "0": "study_end_censor", "0*": "non_tumor_endpoint"}[code],
                "paper_event": int(code == "1"), "composite_event": int(code != "0"),
                "source_sheet": "Figure 6B", "time_cell": f"A{row}", "status_cell": cell.coordinate,
            })
    assert len(records) == 37
    assert Counter(r["source_code"] for r in records) == {"1": 27, "0*": 8, "0": 2}
    assert Counter(r["group"] for r in records) == {g["id"]: g["expected_n"] for g in GROUPS}

    modes = {}
    rmst_rows = []
    for mode in ["paper_compatible", "composite_endpoint"]:
        curves, estimates = {}, []
        for group in GROUPS:
            subset = [r for r in records if r["group"] == group["id"]]
            curves[group["id"]] = km(subset, mode)
            for tau in HORIZONS:
                loo = [rmst(subset[:i] + subset[i + 1:], mode, tau) for i in range(len(subset))]
                supported = [v for v in loo if v is not None]
                point = {
                    "group": group["id"], "mode": mode, "horizon": tau,
                    "rmst_days": rmst(subset, mode, tau),
                    "at_risk_before_horizon": sum(r["time"] >= tau for r in subset),
                    "followed_beyond_horizon": sum(r["time"] > tau for r in subset),
                    "loo_min_days": min(supported) if supported else None,
                    "loo_max_days": max(supported) if supported else None,
                    "loo_supported": len(supported), "loo_total": len(subset),
                }
                estimates.append(point)
                rmst_rows.append(point)
        modes[mode] = {"curves": curves, "rmst": estimates}

    times = [source_bli.cell(row, 1).value for row in range(2, 9)]
    assert times == [3, 6, 10, 13, 17, 21, 26]
    bli_records, series = [], []
    for group in GROUPS:
        assert clean(source_bli.cell(1, group["bli_start"]).value) == group["label"]
        for col in range(group["bli_start"], group["bli_end"] + 1):
            column = openpyxl.utils.get_column_letter(col)
            baseline = source_bli.cell(2, col).value
            assert is_number(baseline) and baseline > 0
            series_id = f"fig6CG-{column}"
            observations, last_terminal = [], None
            for row, time in enumerate(times, start=2):
                cell = source_bli.cell(row, col)
                raw = cell.value
                if is_number(raw):
                    assert raw > 0
                    assert last_terminal is None, (series_id, "numeric measurement after recorded terminal status")
                    status, radiance, fold = "observed", raw, raw / baseline
                elif raw is None:
                    status = {"deceased": "previously_deceased_blank", "censor": "previously_censored_blank"}.get(last_terminal, "missing_unspecified")
                    radiance, fold = None, None
                else:
                    status = clean(raw).lower()
                    assert status in {"deceased", "censor"}, (cell.coordinate, raw)
                    last_terminal, radiance, fold = status, None, None
                point = {
                    "series_id": series_id, "group": group["id"], "time": time,
                    "raw_value": raw, "status": status, "radiance": radiance,
                    "baseline_radiance": baseline, "fold_from_day3": fold,
                    "log2_fold_from_day3": math.log2(fold) if fold is not None else None,
                    "source_sheet": "Figure 6C-G", "source_cell": cell.coordinate,
                }
                observations.append(point)
                bli_records.append(point)
            series.append({"id": series_id, "group": group["id"], "source_column": column, "baseline_radiance": baseline, "observations": observations})
    summaries = []
    for group in GROUPS:
        for time in times:
            points = [r for r in bli_records if r["group"] == group["id"] and r["time"] == time]
            observed = [r for r in points if r["status"] == "observed"]
            folds = [r["fold_from_day3"] for r in observed]
            radiances = [r["radiance"] for r in observed]
            summaries.append({
                "group": group["id"], "time": time, "baseline_n": len(points),
                "observed_n": len(observed), "missing_n": len(points) - len(observed),
                "missing_status_counts": dict(Counter(r["status"] for r in points if r["status"] != "observed")),
                "median_fold": statistics.median(folds) if folds else None,
                "min_fold": min(folds) if folds else None, "max_fold": max(folds) if folds else None,
                "median_radiance": statistics.median(radiances) if radiances else None,
                "mean_radiance": statistics.mean(radiances) if radiances else None,
                "complete_baseline_cohort": len(observed) == len(points),
            })
    numeric_n = sum(r["status"] == "observed" for r in bli_records)
    assert len(series) == 37 and len(bli_records) == 259 and numeric_n == 238
    assert Counter(r["status"] for r in bli_records) == {"observed": 238, "deceased": 15, "censor": 2, "previously_deceased_blank": 4}
    assert len([r for r in records if r["group"] == "rli_tmz" and r["source_code"] == "0*"]) == 7
    common_complete = [t for t in times if all(s["complete_baseline_cohort"] for s in summaries if s["time"] == t)]
    assert common_complete == [3, 6, 10, 13, 17]

    provenance = {
        "article": "Haddad et al. (2026), Non-lytic viral immunotherapy induces long-term glioblastoma survival and tumor-specific immunity without eliciting an antiviral response",
        "doi": "10.1038/s41467-026-72746-5", "article_url": ARTICLE_URL,
        "source_url": SOURCE_URL, "source_filename": SOURCE_NAME, "sha256": SOURCE_SHA256,
        "source_license": "CC BY 4.0", "license_url": "https://creativecommons.org/licenses/by/4.0/",
        "source_bytes": args.source.stat().st_size, "source_sheets": ["Figure 6B", "Figure 6C-G"],
        "accessed": "2026-09-30", "biological_unit": "mouse within one figure; no cross-sheet animal identifiers",
        "code_origin": "New independent descriptive analysis; the source publication states that no custom code was developed or used.",
        "figure_legend_facts": [
            "Figure 6 survival compares PBS+Vehicle (10), PBS+TMZ (9), RRV-RLI+Vehicle (9), RRV-RLI+TMZ (9).",
            "The Figure 6 legend identifies seven combination-group non-tumor endpoints as TMZ toxicity; the vehicle-group non-tumor endpoint cause is not assigned here.",
            "Figure 6D-G represents individual mouse bioluminescent trajectories; time is days after tumor implantation.",
            "The workbook labels radiance ph/sec/ROI; no conversion to tumor size or cell number is performed.",
            "Pretreatment BLI was used for treatment-group randomization; the workbook supplies no raw imaging files or camera metadata.",
        ],
    }
    limitations = [
        "Exploratory reanalysis of published small-cohort mouse data, not a new experiment, clinical guidance, or confirmation of therapeutic efficacy.",
        "Paper-compatible analysis censors non-tumor endpoints. The composite counts recorded tumor and non-tumor endpoints: a different estimand, not a correction to the paper or an all-cause mortality estimate.",
        "Non-tumor censoring can be informative. This sensitivity exposes the dependence on endpoint definition but does not solve informative censoring or estimate a treatment-policy effect.",
        "Workbook columns identify longitudinal series only within Figure 6C-G. Survival rows cannot be matched to BLI columns; no joint prediction, mediation, or subject-level BLI-survival correlation is justified.",
        "BLI summaries condition on observed, baseline-valid measurements. Death/censor/blank is never converted to zero; later medians do not summarize the original randomized cohort.",
        "Last fully observed common BLI date is day 17; later comparisons change the observed denominator.",
        "Default RMST horizon 60 days was selected before computing results. Other horizons are exploratory; the long tail is often supported by one mouse.",
        "Leave-one-out extrema are influence checks, not confidence intervals. At longer horizons some deletions leave positive unsupported tails and are reported as unestimable.",
        "The data do not establish optimal dosing, tissue collection schedules, true tumor eradication, or a cause for an unlabeled non-tumor endpoint.",
    ]
    result = {
        "schema_version": "1.0", "provenance": provenance,
        "study": {"model": "SB28 murine glioblastoma", "time_unit": "days after tumor implantation", "radiance_unit": "ph/sec/ROI as supplied", "default_horizon": 60},
        "groups": GROUPS,
        "survival": {"records": records, "modes": modes, "horizons": HORIZONS, "counts": dict(Counter(r["source_code"] for r in records))},
        "bli": {"times": times, "series": series, "records": bli_records, "summaries": summaries, "qc": {"baseline_series_n": len(series), "measurement_slots": len(bli_records), "numeric_measurements": numeric_n, "status_counts": dict(Counter(r["status"] for r in bli_records)), "common_complete_days": common_complete, "last_common_complete_day": max(common_complete)}},
        "limitations": limitations,
    }
    data_dir = ROOT / "data"
    data_dir.mkdir(parents=True, exist_ok=True)
    json_write(data_dir / "data.json", result)
    json_write(data_dir / "provenance.json", provenance)
    csv_write(data_dir / "survival_records.csv", records, list(records[0]))
    csv_write(data_dir / "bli_records.csv", bli_records, list(bli_records[0]))
    csv_write(data_dir / "rmst_estimates.csv", rmst_rows, list(rmst_rows[0]))
    summary_csv = [{**s, "missing_status_counts": json.dumps(s["missing_status_counts"], sort_keys=True)} for s in summaries]
    csv_write(data_dir / "bli_summaries.csv", summary_csv, list(summaries[0]))
    print(json.dumps({"survival_n": len(records), "survival_codes": result["survival"]["counts"], "bli_qc": result["bli"]["qc"], "rmst_60": [r for r in rmst_rows if r["horizon"] == 60]}, indent=2))


if __name__ == "__main__":
    main()
