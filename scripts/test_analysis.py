#!/usr/bin/env python3
"""Independent exact-arithmetic/source checks; run after build_data.py."""
import csv
import importlib.util
import json
import math
from fractions import Fraction
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
data = json.loads((ROOT / "data/data.json").read_text())
workbook = openpyxl.load_workbook(ROOT / "research-inputs/41467_2026_72746_MOESM4_ESM.xlsx", data_only=True)
sheet = workbook["Figure 6B"]
checked = 0
for mode, payload in data["survival"]["modes"].items():
    for group in data["groups"]:
        column = group["survival_column"]
        # Independently read the original spreadsheet, not the parsed records.
        source = [(sheet[f"A{row}"].value, str(sheet[f"{column}{row}"].value)) for row in range(2, sheet.max_row + 1) if sheet[f"{column}{row}"].value is not None]
        events = {"1"} if mode == "paper_compatible" else {"1", "0*"}
        event_dates = sorted({t for t, code in source if code in events})
        for estimate in [x for x in payload["rmst"] if x["group"] == group["id"]]:
            tau = estimate["horizon"]
            # Integrate only event-date breakpoints with exact rational numbers.
            area, survival, previous = Fraction(0), Fraction(1), 0
            for t in [x for x in event_dates if x < tau]:
                area += (t - previous) * survival
                risk = sum(time >= t for time, _ in source)
                failures = sum(time == t and code in events for time, code in source)
                survival *= Fraction(risk - failures, risk)
                previous = t
            area += (tau - previous) * survival
            assert math.isclose(float(area), estimate["rmst_days"], abs_tol=1e-10)
            if mode == "composite_endpoint":
                # All residual censors are at day399; truncated empirical means
                # therefore independently equal RMST for all included horizons.
                empirical = sum(min(t, tau) for t, _ in source) / len(source)
                assert math.isclose(empirical, estimate["rmst_days"], abs_tol=1e-10)
            checked += 1
        for point in payload["curves"][group["id"]][1:]:
            probability = Fraction(1)
            for t in [x for x in event_dates if x <= point["time"]]:
                risk = sum(time >= t for time, _ in source)
                failures = sum(time == t and code in events for time, code in source)
                probability *= Fraction(risk - failures, risk)
            assert math.isclose(float(probability), point["survival"], abs_tol=1e-12)

bli = workbook["Figure 6C-G"]
for record in data["bli"]["records"]:
    raw = bli[record["source_cell"]].value
    assert raw == record["raw_value"]
    if record["radiance"] is None:
        assert record["fold_from_day3"] is None and record["log2_fold_from_day3"] is None
    else:
        column = bli[record["source_cell"]].column
        expected = raw / bli.cell(2, column).value
        assert math.isclose(expected, record["fold_from_day3"], abs_tol=1e-12)

for filename, expected in [("survival_records.csv", 37), ("bli_records.csv", 259), ("bli_summaries.csv", 28), ("rmst_estimates.csv", 40)]:
    with (ROOT / "data" / filename).open() as handle:
        assert len(list(csv.DictReader(handle))) == expected

spec = importlib.util.spec_from_file_location("builder", ROOT / "scripts/build_data.py")
builder = importlib.util.module_from_spec(spec)
spec.loader.exec_module(builder)
synthetic = [{"time": 2, "paper_event": 1}, {"time": 2, "paper_event": 0}, {"time": 4, "paper_event": 0}]
assert math.isclose(builder.km(synthetic, "paper_compatible")[1]["survival"], 2 / 3)
assert builder.rmst(synthetic, "paper_compatible", 5) is None
print(f"PASS: {checked} exact-arithmetic RMST checks; all KM points; 259 BLI cell/normalization checks; CSV counts; tied-event and unsupported-tail tests.")
