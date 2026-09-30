#!/usr/bin/env python3
"""Independently verify Figure 6 outputs without importing the analysis builder.

Requires only openpyxl. Reads source workbook cells directly. Kaplan-Meier
probabilities use exact fractions and direct risk-set counts. RMST is calculated
from event-probability masses, not by copying a step-area integration routine.
No source data or website files are modified. --report optionally writes JSON.
"""

import argparse
from collections import Counter
from fractions import Fraction
import hashlib
from io import BytesIO
import json
import math
from pathlib import Path
from statistics import mean, median
from urllib.request import urlopen

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
FILENAME = "41467_2026_72746_MOESM4_ESM.xlsx"
SOURCE_URL = (
    "https://media.springernature.com/original/springer-static/esm/"
    "art%3A10.1038%2Fs41467-026-72746-5/MediaObjects/" + FILENAME
)
SOURCE_SHA = "3e302d90157981d46c04b72bc94633ad2fdf55381f016a33455d3986f35589e6"
GROUPS = (
    ("pbs_vehicle", "PBS + Vehicle", 3, 2, 11, 10),
    ("pbs_tmz", "PBS + TMZ", 4, 12, 20, 9),
    ("rli_vehicle", "RRV-RLI + Vehicle", 5, 21, 29, 9),
    ("rli_tmz", "RRV-RLI + TMZ", 6, 30, 38, 9),
)
TIMES = [3, 6, 10, 13, 17, 21, 26]
HORIZONS = [20, 40, 60, 120, 399]
MODES = {"paper_compatible": {"1"}, "composite_endpoint": {"1", "0*"}}


def equal(actual, expected, context):
    if isinstance(expected, (float, Fraction)):
        if actual is None or not math.isclose(
            float(actual), float(expected), rel_tol=1e-11, abs_tol=1e-10
        ):
            raise AssertionError(f"{context}: {actual!r} != {expected!r}")
    elif actual != expected:
        raise AssertionError(f"{context}: {actual!r} != {expected!r}")


def fields(actual, expected, context):
    for key, value in expected.items():
        equal(actual[key], value, f"{context}.{key}")


def numbered(rows, key):
    result = {row[key]: row for row in rows}
    equal(len(result), len(rows), f"unique {key}")
    return result


def parse_source(workbook):
    endpoints, series, cells, summaries = [], [], [], []
    survival_sheet = workbook["Figure 6B"]
    imaging_sheet = workbook["Figure 6C-G"]
    equal([imaging_sheet.cell(row, 1).value for row in range(2, 9)], TIMES, "day grid")
    for group, label, endpoint_col, first, last, n in GROUPS:
        equal(survival_sheet.cell(1, endpoint_col).value.replace("\xa0", " "), label, "survival header")
        equal(imaging_sheet.cell(1, first).value.replace("\xa0", " "), label, "imaging header")
        endpoint_start = len(endpoints)
        for row in range(2, survival_sheet.max_row + 1):
            code = survival_sheet.cell(row, endpoint_col).value
            if code is None:
                continue
            code = str(code)
            if code not in {"1", "0", "0*"}:
                raise AssertionError(f"Unexpected endpoint code: {code}")
            endpoints.append({
                "id": f"fig6B-r{row}", "group": group,
                "time": survival_sheet.cell(row, 1).value, "source_code": code,
                "event_type": {"1": "tumor_endpoint", "0": "study_end_censor", "0*": "non_tumor_endpoint"}[code],
                "paper_event": int(code == "1"), "composite_event": int(code != "0"),
                "source_sheet": "Figure 6B", "time_cell": f"A{row}",
                "status_cell": f"{get_column_letter(endpoint_col)}{row}",
            })
        equal(len(endpoints) - endpoint_start, n, f"{group} endpoint n")
        for col in range(first, last + 1):
            letter = get_column_letter(col)
            baseline = imaging_sheet.cell(2, col).value
            if not isinstance(baseline, (int, float)) or baseline <= 0:
                raise AssertionError("Invalid baseline")
            entries, previous_terminal = [], None
            for row, day in zip(range(2, 9), TIMES):
                raw = imaging_sheet.cell(row, col).value
                if isinstance(raw, (int, float)) and not isinstance(raw, bool):
                    status, radiance = "observed", raw
                    fold, logfold = raw / baseline, math.log2(raw / baseline)
                else:
                    radiance = fold = logfold = None
                    if raw in ("deceased", "censor"):
                        status, previous_terminal = raw, raw
                    elif raw is None:
                        status = f"previously_{previous_terminal}_blank" if previous_terminal else "blank"
                    else:
                        raise AssertionError(f"Unexpected BLI annotation: {raw!r}")
                entries.append({
                    "series_id": f"fig6CG-{letter}", "group": group, "time": day,
                    "raw_value": raw, "status": status, "radiance": radiance,
                    "baseline_radiance": baseline, "fold_from_day3": fold,
                    "log2_fold_from_day3": logfold, "source_sheet": "Figure 6C-G",
                    "source_cell": f"{letter}{row}",
                })
            cells.extend(entries)
            series.append({"id": f"fig6CG-{letter}", "group": group, "source_column": letter,
                           "baseline_radiance": baseline, "observations": entries})
        for day in TIMES:
            rows = [r for r in cells if r["group"] == group and r["time"] == day]
            observed = [r for r in rows if r["status"] == "observed"]
            folds = [r["fold_from_day3"] for r in observed]
            radiances = [r["radiance"] for r in observed]
            summaries.append({
                "group": group, "time": day, "baseline_n": n,
                "observed_n": len(observed), "missing_n": n - len(observed),
                "missing_status_counts": dict(Counter(r["status"] for r in rows if r["status"] != "observed")),
                "median_fold": median(folds) if folds else None,
                "min_fold": min(folds) if folds else None, "max_fold": max(folds) if folds else None,
                "median_radiance": median(radiances) if radiances else None,
                "mean_radiance": mean(radiances) if radiances else None,
                "complete_baseline_cohort": len(observed) == n,
            })
    return endpoints, series, cells, summaries


def event_atoms(records, event_codes):
    """Product-limit event masses from direct >= time risk-set counts."""
    event_times = sorted({r["time"] for r in records if r["source_code"] in event_codes})
    atoms, survival = [], Fraction(1)
    for time in event_times:
        risk = sum(r["time"] >= time for r in records)
        events = sum(r["time"] == time and r["source_code"] in event_codes for r in records)
        mass = survival * Fraction(events, risk)
        atoms.append((time, mass))
        survival -= mass
    return atoms, survival


def km_curve(records, event_codes):
    atoms, _ = event_atoms(records, event_codes)
    result = []
    for time in [0] + sorted({r["time"] for r in records}):
        events = sum(r["time"] == time and r["source_code"] in event_codes for r in records)
        censored = sum(r["time"] == time and r["source_code"] not in event_codes for r in records)
        result.append({
            "time": time, "survival": Fraction(1) - sum((mass for t, mass in atoms if t <= time), Fraction(0)),
            "at_risk": sum(r["time"] >= time for r in records),
            "events": events, "censored": censored,
            "remaining_after": sum(r["time"] > time for r in records),
        })
    return result


def restricted_mean(records, event_codes, horizon):
    """E[min(T,tau)] = sum(t * KM event mass) + tau * S(tau)."""
    atoms, final_survival = event_atoms(records, event_codes)
    if horizon > max(r["time"] for r in records) and final_survival > 0:
        return None
    relevant = [(t, mass) for t, mass in atoms if t <= horizon]
    s_tau = Fraction(1) - sum((mass for _, mass in relevant), Fraction(0))
    return sum((Fraction(t) * mass for t, mass in relevant), Fraction(0)) + horizon * s_tau


def verify(source_bytes, derived):
    equal(hashlib.sha256(source_bytes).hexdigest(), SOURCE_SHA, "source checksum")
    endpoints, series, cells, summaries = parse_source(load_workbook(BytesIO(source_bytes), data_only=True))
    equal(len(endpoints), 37, "endpoint records")
    equal(len(cells), 259, "BLI slots")
    equal(len(derived["groups"]), 4, "group metadata length")
    for actual, (group, label, col, first, last, n) in zip(derived["groups"], GROUPS):
        fields(actual, {"id": group, "label": label, "survival_column": get_column_letter(col),
                        "bli_start": first, "bli_end": last, "expected_n": n}, "group metadata")
    endpoint_data = numbered(derived["survival"]["records"], "id")
    equal(set(endpoint_data), {r["id"] for r in endpoints}, "endpoint keys")
    for record in endpoints:
        fields(endpoint_data[record["id"]], record, record["id"])
    cell_data = numbered(derived["bli"]["records"], "source_cell")
    equal(set(cell_data), {r["source_cell"] for r in cells}, "BLI cell keys")
    for record in cells:
        fields(cell_data[record["source_cell"]], record, record["source_cell"])
    series_data = numbered(derived["bli"]["series"], "id")
    equal(set(series_data), {s["id"] for s in series}, "series keys")
    for item in series:
        actual = series_data[item["id"]]
        fields(actual, {k: v for k, v in item.items() if k != "observations"}, item["id"])
        equal(len(actual["observations"]), 7, "within-series length")
        for a, b in zip(actual["observations"], item["observations"]):
            fields(a, b, item["id"])
    summary_data = {(r["group"], r["time"]): r for r in derived["bli"]["summaries"]}
    equal(len(derived["bli"]["summaries"]), 28, "summary rows")
    equal(len(summary_data), 28, "summary keys")
    for item in summaries:
        fields(summary_data[(item["group"], item["time"])], item, f"summary {item['group']} {item['time']}")
    counts = dict(Counter(r["source_code"] for r in endpoints))
    equal(derived["survival"]["counts"], counts, "endpoint counts")
    equal(derived["survival"]["horizons"], HORIZONS, "RMST horizons")
    statuses = dict(Counter(r["status"] for r in cells))
    complete_days = [day for day in TIMES if all(r["status"] == "observed" for r in cells if r["time"] == day)]
    fields(derived["bli"]["qc"], {
        "baseline_series_n": 37, "measurement_slots": 259,
        "numeric_measurements": statuses["observed"], "status_counts": statuses,
        "common_complete_days": complete_days, "last_common_complete_day": max(complete_days),
    }, "BLI QC")
    equal(derived["bli"]["times"], TIMES, "BLI times")

    rmst_results, curve_count, deletion_count = [], 0, 0
    equal(set(derived["survival"]["modes"]), set(MODES), "endpoint modes")
    for mode, codes in MODES.items():
        actual_mode = derived["survival"]["modes"][mode]
        equal(set(actual_mode["curves"]), {g[0] for g in GROUPS}, "curve groups")
        equal(len(actual_mode["rmst"]), 20, f"{mode} RMST row count")
        rmst_data = {(r["group"], r["horizon"]): r for r in actual_mode["rmst"]}
        equal(len(rmst_data), 20, f"{mode} RMST rows")
        for group, _, _, _, _, n in GROUPS:
            group_rows = [r for r in endpoints if r["group"] == group]
            curve = km_curve(group_rows, codes)
            equal(len(actual_mode["curves"][group]), len(curve), "curve length")
            for a, b in zip(actual_mode["curves"][group], curve):
                fields(a, b, f"KM {mode} {group} {b['time']}")
                curve_count += 1
            for horizon in HORIZONS:
                estimate = restricted_mean(group_rows, codes, horizon)
                loo = [restricted_mean(group_rows[:i] + group_rows[i + 1:], codes, horizon) for i in range(n)]
                deletion_count += n
                supported = [v for v in loo if v is not None]
                expected = {
                    "group": group, "mode": mode, "horizon": horizon, "rmst_days": estimate,
                    "at_risk_before_horizon": sum(r["time"] >= horizon for r in group_rows),
                    "followed_beyond_horizon": sum(r["time"] > horizon for r in group_rows),
                    "loo_min_days": min(supported) if supported else None,
                    "loo_max_days": max(supported) if supported else None,
                    "loo_supported": len(supported), "loo_total": n,
                }
                fields(rmst_data[(group, horizon)], expected, f"RMST {mode} {group} {horizon}")
                rmst_results.append({k: float(v) if isinstance(v, Fraction) else v for k, v in expected.items()})
    delta = {}
    for mode in MODES:
        selected = {r["group"]: r["rmst_days"] for r in rmst_results if r["mode"] == mode and r["horizon"] == 60}
        delta[mode] = selected["rli_tmz"] - selected["rli_vehicle"]
    equal(delta["paper_compatible"], Fraction(383, 30), "default paper-compatible contrast")
    equal(delta["composite_endpoint"], 3.0, "default composite contrast")
    return {
        "status": "PASS", "source_sha256": SOURCE_SHA,
        "independence": "No builder imports or code reuse. Source-cell parsing plus exact rational event-mass RMST and direct KM risk sets.",
        "verified": {"endpoint_records": 37, "bli_cells": 259, "bli_series": 37, "bli_summaries": 28,
                     "km_curve_points": curve_count, "rmst_estimates": len(rmst_results), "leave_one_out_recomputations": deletion_count},
        "endpoint_code_counts": counts, "bli_status_counts": statuses,
        "last_common_complete_imaging_day": max(complete_days),
        "default_60_day_rli_tmz_minus_rli_vehicle": delta,
        "all_observed_medians": summaries, "all_rmst_results": rmst_results,
        "boundaries": [
            "259 imaging slots are repeated measures from 37 within-sheet series, not 259 independent animals.",
            "No animal identifier connects the endpoint and imaging sheets; no cross-sheet pairing or predictive association is validated.",
            "Non-tumor endpoints change the estimand when counted. The composite is not established all-cause mortality.",
            "Observed-only medians can be affected by attrition; missing and terminal annotations are not zero radiance.",
            "Leave-one-record-out ranges are influence diagnostics, not confidence intervals. Unsupported positive-tail deletions are omitted explicitly.",
            "Censoring assumptions, small groups and single-animal late risk sets preclude causal clinical claims or a standalone treatment ranking.",
        ],
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path)
    parser.add_argument("--derived", type=Path, default=ROOT / "data" / "data.json")
    parser.add_argument("--report", type=Path, help="Optional detailed QA JSON; source and website data remain unchanged.")
    args = parser.parse_args()
    local = args.source or ROOT / "research-inputs" / FILENAME
    if local.exists():
        source_bytes = local.read_bytes()
    elif args.source:
        raise FileNotFoundError(local)
    else:
        with urlopen(SOURCE_URL, timeout=60) as response:
            source_bytes = response.read()
    derived_bytes = args.derived.read_bytes()
    report = verify(source_bytes, json.loads(derived_bytes))
    report["derived_sha256"] = hashlib.sha256(derived_bytes).hexdigest()
    if args.report:
        args.report.write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({k: report[k] for k in ("status", "verified", "default_60_day_rli_tmz_minus_rli_vehicle")}))


if __name__ == "__main__":
    main()
